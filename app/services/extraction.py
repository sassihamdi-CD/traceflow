"""Background extraction, branched by file type (Correction 1).

- PDF: file bytes go DIRECTLY to Claude as a native document input.
  No pdfplumber/pypdf pre-extraction. Claude's own PDF reading handles
  layout/tables; the `location` field uses Claude's page reference.
- XLSX: parsed with openpyxl. `location` is the REAL sheet/cell coordinate
  (e.g. "Sheet 1 / F8") taken from openpyxl, never inferred by the model.
  Cell values are flattened into a structured table; Claude only maps that
  table onto field_definitions.
- CSV: same path as XLSX via pandas (flat rows, "Row N" locations).
- Unknown types never reach here: the upload endpoint rejects them (415).

Every inserted row uses status='proposed', extracted_by='ai_system'.
Extraction NEVER writes accepted/corrected (spec section 1).
"""
from __future__ import annotations

import base64
import io
import json
import re

import pandas as pd
from anthropic import AsyncAnthropic
from openpyxl import load_workbook

from app.config import settings
from app.services.extraction_prompts import (
    SYSTEM_PROMPT,
    pdf_task_instruction,
    table_task_instruction,
)


class TruncatedList(list):
    """A plain list that remembers whether a pipeline cap cut it short.

    Behaves exactly like a list (iteration, indexing, len, json.dumps all
    work), so existing call sites are unaffected; new code reads `.truncated`
    to surface the flag on the document row + UI badge.
    """

    def __init__(self, items=(), truncated: bool = False):
        super().__init__(items)
        self.truncated = truncated


# Row caps feeding the truncation flag (mirrors prompt + doc budgets).
TABLE_MAX_ROWS = 500
PARTIES_CAP = 12

# PDF evidence pins must look like Claude's page reference, e.g. 'Page 3'.
PDF_LOCATION_RE = re.compile(r"^Page \d+(?:-\d+)?$")


def table_locations(table: list[dict] | None) -> set[str]:
    """Every location string the model is allowed to echo back."""
    valid: set[str] = set()
    for r in table or []:
        if isinstance(r, dict) and r.get("sheet") and r.get("cell"):
            valid.add(f"{r['sheet']} / {r['cell']}")
    return valid


def coerce_pdf_location(location: object) -> str | None:
    """Keep Claude's page reference only when it matches 'Page N'."""
    if not isinstance(location, str):
        return None
    loc = re.sub(r"\s+", " ", location.strip())
    return loc if PDF_LOCATION_RE.match(loc) else None


def coerce_table_location(location: object, valid: set[str]) -> str | None:
    """Keep a sheet/cell pin only when it exists in the parsed table."""
    if not isinstance(location, str):
        return None
    loc = location.strip()
    return loc if loc in valid else None


def validate_item_locations(items: list[dict], ext: str, table: list[dict] | None = None) -> list[dict]:
    """Post-hoc coordinate check on extracted passport values.

    The model is told to copy coordinates verbatim, but it still invents
    them: XLSX/CSV locations must exist in the parsed table, PDF locations
    must match 'Page N'. Unprovable pins are nulled (the VALUE is kept for
    the reviewer) — a fabricated pin is worse than a missing one.
    Never raises; unknown ext keeps locations untouched.
    """
    out: list[dict] = []
    valid = table_locations(table) if ext in (".xlsx", ".csv") else set()
    for it in items:
        it = dict(it)
        loc = it.get("location")
        if ext in (".xlsx", ".csv"):
            it["location"] = coerce_table_location(loc, valid)
        elif ext == ".pdf":
            it["location"] = coerce_pdf_location(loc)
        out.append(it)
    return out


def xlsx_to_table(xlsx_bytes: bytes, max_rows: int = TABLE_MAX_ROWS) -> TruncatedList:
    """Flatten workbook cells to [{sheet, cell, value}]. Coordinates from openpyxl."""
    wb = load_workbook(filename=io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    rows = TruncatedList()
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=False):
                for cell in row:
                    if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
                        continue
                    if len(rows) >= max_rows:
                        # A (max+1)th non-empty cell exists: the cap bit off data.
                        rows.truncated = True
                        return rows
                    rows.append({"sheet": ws.title, "cell": cell.coordinate, "value": str(cell.value)})
    finally:
        wb.close()
    return rows


def csv_to_table(csv_bytes: bytes, max_rows: int = TABLE_MAX_ROWS) -> TruncatedList:
    df = pd.read_csv(io.BytesIO(csv_bytes), dtype=str, keep_default_na=False)
    rows = TruncatedList()
    for i, record in enumerate(df.to_dict(orient="records"), start=2):  # row 1 = header
        for col, val in record.items():
            if val is None or (isinstance(val, str) and not val.strip()):
                continue
            if len(rows) >= max_rows:
                rows.truncated = True
                return rows
            rows.append({"sheet": "CSV", "cell": f"Row {i} / {col}", "value": str(val)})
    return rows


def build_mapping_prompt(field_defs: list[dict], table: list[dict]) -> str:
    allowed = ", ".join(f"{d['field_key']} ({d['label']})" for d in field_defs)
    return table_task_instruction(allowed, json.dumps(table[:500]))


def sanitize_items(items: object, allowed: set[str]) -> list[dict]:
    """Drop anything the model returned that is not a well-formed hit.

    Only well-formed rows survive: known field_key, non-empty value.
    Runs before DB insert; extraction inserts proposed rows only.
    """
    clean: list[dict] = []
    if not isinstance(items, list):
        return clean
    for it in items:
        if not isinstance(it, dict):
            continue
        key = it.get("field_key")
        val = it.get("value")
        if key not in allowed or not isinstance(val, str) or not val.strip():
            continue
        clean.append({
            "field_key": key,
            "value": val.strip(),
            "unit": it.get("unit") if isinstance(it.get("unit"), str) else None,
            "location": it.get("location") if isinstance(it.get("location"), str) else None,
        })
    return clean


class ExtractionParseError(Exception):
    """Claude's reply contained no parseable JSON payload."""


def parse_extraction_payload(text: str) -> list[dict]:
    """Tolerant parser for the model's JSON reply.

    Handles: raw array, {"items": [...]}, markdown-fenced JSON, and leading /
    trailing prose. Raises ExtractionParseError when no JSON array is found.
    """
    candidate = text.strip()
    if candidate.startswith("```"):
        # strip ```json ... ``` fences
        lines = candidate.splitlines()
        lines = [ln for ln in lines if not ln.strip().startswith("```")]
        candidate = "\n".join(lines).strip()
    try:
        data = json.loads(candidate)
    except json.JSONDecodeError:
        # fall back: first [...] block in the reply
        start, end = candidate.find("["), candidate.rfind("]")
        if start == -1 or end <= start:
            raise ExtractionParseError("no JSON array in model reply")
        try:
            data = json.loads(candidate[start : end + 1])
        except json.JSONDecodeError as e:
            raise ExtractionParseError(f"malformed JSON array: {e}")
    if isinstance(data, dict):
        data = data.get("items", [])
    if not isinstance(data, list):
        raise ExtractionParseError("payload is not a list")
    return data


async def _claude_json_messages(messages: list[dict]) -> list[dict]:
    client = AsyncAnthropic(api_key=settings.anthropic_api_key)
    resp = await client.messages.create(
        model=settings.anthropic_model,
        max_tokens=2000,
        temperature=0,
        system=SYSTEM_PROMPT,
        messages=messages,
    )
    text = "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
    return parse_extraction_payload(text)


async def extract_pdf_native(pdf_bytes: bytes, field_defs: list[dict]) -> list[dict]:
    allowed = ", ".join(f"{d['field_key']} ({d['label']})" for d in field_defs)
    b64 = base64.standard_b64encode(pdf_bytes).decode("ascii")
    return await _claude_json_messages(
        [
            {
                "role": "user",
                "content": [
                    {
                        "type": "document",
                        "source": {"type": "base64", "media_type": "application/pdf", "data": b64},
                    },
                    {"type": "text", "text": pdf_task_instruction(allowed)},
                ],
            }
        ]
    )


async def extract_table_mapped(table: list[dict], field_defs: list[dict]) -> list[dict]:
    return await _claude_json_messages(
        [{"role": "user", "content": [{"type": "text", "text": build_mapping_prompt(field_defs, table)}]}]
    )


PARTIES_SYSTEM = (
    "You list the organizations named in a supplier document. "
    "Output STRICT JSON only: an array of objects with exactly the keys "
    "name and role. Role is one of: applicant, vendor, supplier, laboratory, "
    "brand, other. List each distinct organization once with its clearest role. "
    "If none are named, output []. No prose, no markdown."
)


def cap_parties(data: object) -> TruncatedList:
    """Pure parties sanitizer: [{name, role}], capped at PARTIES_CAP.

    Split out of extract_parties_pdf so the 12-cap + truncation flag are
    unit-testable without network. Never raises.
    """
    clean = TruncatedList()
    try:
        if not isinstance(data, list):
            return clean
        kept: list[dict] = []
        for it in data:
            if isinstance(it, dict) and isinstance(it.get("name"), str) and it["name"].strip():
                role = it.get("role") if it.get("role") in (
                    "applicant", "vendor", "supplier", "laboratory", "brand", "other") else "other"
                kept.append({"name": it["name"].strip()[:120], "role": role})
        clean.extend(kept[:PARTIES_CAP])
        clean.truncated = len(kept) > PARTIES_CAP
        return clean
    except Exception:
        return clean


async def extract_parties_pdf(pdf_bytes: bytes) -> list[dict]:
    """Who is on this document? Failure-tolerant: [] on any problem."""
    try:
        client = AsyncAnthropic(api_key=settings.anthropic_api_key)
        b64 = base64.standard_b64encode(pdf_bytes).decode("ascii")
        resp = await client.messages.create(
            model=settings.anthropic_model,
            max_tokens=400,
            temperature=0,
            system=PARTIES_SYSTEM,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "document", "source": {
                        "type": "base64", "media_type": "application/pdf", "data": b64}},
                    {"type": "text", "text": "List the organizations named in this document with their roles. JSON array only."},
                ],
            }],
        )
        text = "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
        data = parse_extraction_payload(text)
        return cap_parties(data)
    except Exception:
        return []
